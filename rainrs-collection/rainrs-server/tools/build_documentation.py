"""Generate the rainrs-server documentation (``.docx``) and its diagrams.

The script renders the UML/ERD diagrams with matplotlib and assembles
them into a Word document with python-docx.

The data model diagram and the data model tables are generated from the
catalogue in :mod:`rainrs_server.model.schema`, so the documentation
always describes the data model that the code actually has.

Usage (from the ``rainrs-server`` directory)::

    uv run python tools/build_documentation.py

The output is written to ``docs/rainrs-server-documentation.docx``. The
diagram images are rendered into a temporary directory and embedded in
the document; no image files are left in the repository.
"""

from __future__ import annotations

import tempfile
from pathlib import Path

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt  # noqa: E402 - must follow the backend selection
from docx import Document  # noqa: E402 - documentation tooling
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Inches, Pt  # noqa: E402
from matplotlib.patches import FancyBboxPatch, Polygon  # noqa: E402

from rainrs_server.model import schema as data_model  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "rainrs-server-documentation.docx"

BOX_FILL = "#eef4fb"
BOX_EDGE = "#3b6ea5"
DECISION_FILL = "#fff4e5"
DECISION_EDGE = "#d08b2c"
GROUP_FILL = "#f4f6f8"
GROUP_EDGE = "#9aa5b1"
ACTOR_FILL = "#e8f4ea"
ACTOR_EDGE = "#4c8c5a"
STORE_FILL = "#f3e9f7"
STORE_EDGE = "#8a5aa5"

LINE_WIDTH = 1.4
TITLE_OFFSET = 0.9
LINE_OFFSET = 1.9
LINE_SPACING = 1.9
MIN_BOX_HEIGHT = 2.4


# ---------------------------------------------------------------------------
# Drawing helpers
# ---------------------------------------------------------------------------
def new_canvas(width: float = 12.0, height: float = 9.0):
    """Return a blank, axis-less canvas for a diagram."""
    figure, axes = plt.subplots(figsize=(width, height))
    axes.set_xlim(0, 100)
    axes.set_ylim(0, 100)
    axes.set_aspect("auto")
    axes.axis("off")
    return figure, axes


def box_height(lines: tuple[str, ...]) -> float:
    """Return the height a box needs for ``lines``."""
    if not lines:
        return MIN_BOX_HEIGHT
    return TITLE_OFFSET + LINE_OFFSET + LINE_SPACING * (len(lines) - 1) + 1.0


def draw_box(
    axes,
    x: float,
    y: float,
    title: str,
    lines: tuple[str, ...] = (),
    *,
    width: float = 16.0,
    height: float | None = None,
    fill: str = BOX_FILL,
    edge: str = BOX_EDGE,
    title_size: float = 9,
    body_size: float = 7.5,
) -> float:
    """Draw a labelled rectangle and return its height.

    The box grows with the number of text lines, so the content always
    stays inside the frame.
    """
    box_height_value = height if height is not None else box_height(lines)
    patch = FancyBboxPatch(
        (x, y),
        width,
        box_height_value,
        boxstyle="round,pad=0.35,rounding_size=0.8",
        linewidth=LINE_WIDTH,
        edgecolor=edge,
        facecolor=fill,
    )
    axes.add_patch(patch)
    axes.text(
        x + width / 2,
        y + box_height_value - TITLE_OFFSET,
        title,
        ha="center",
        va="top",
        fontsize=title_size,
        fontweight="bold",
        color="#1f2933",
    )
    for index, line in enumerate(lines):
        axes.text(
            x + width / 2,
            y + box_height_value - TITLE_OFFSET - LINE_OFFSET - index * LINE_SPACING,
            line,
            ha="center",
            va="top",
            fontsize=body_size,
            color="#33414e",
        )
    return box_height_value


def draw_centered_box(
    axes,
    centre: float,
    y: float,
    title: str,
    lines: tuple[str, ...] = (),
    *,
    width: float = 16.0,
    height: float | None = None,
    fill: str = BOX_FILL,
    edge: str = BOX_EDGE,
) -> float:
    """Draw a box centred on ``centre`` and return its height."""
    return draw_box(
        axes,
        centre - width / 2,
        y,
        title,
        lines,
        width=width,
        height=height,
        fill=fill,
        edge=edge,
    )


def draw_actor(axes, x: float, y: float, label: str, *, width: float = 8.0) -> float:
    """Draw an actor (a rounded box in this diagram style)."""
    height = 3.4
    patch = FancyBboxPatch(
        (x, y),
        width,
        height,
        boxstyle="round,pad=0.35,rounding_size=1.2",
        linewidth=LINE_WIDTH,
        edgecolor=ACTOR_EDGE,
        facecolor=ACTOR_FILL,
    )
    axes.add_patch(patch)
    axes.text(
        x + width / 2,
        y + height / 2,
        label,
        ha="center",
        va="center",
        fontsize=8.5,
        fontweight="bold",
        color="#1f2933",
    )
    return height


def draw_store(axes, x: float, y: float, label: str, *, width: float = 8.0) -> float:
    """Draw a store (rounded box)."""
    height = 3.4
    patch = FancyBboxPatch(
        (x, y),
        width,
        height,
        boxstyle="round,pad=0.35,rounding_size=1.2",
        linewidth=LINE_WIDTH,
        edgecolor=STORE_EDGE,
        facecolor=STORE_FILL,
    )
    axes.add_patch(patch)
    axes.text(
        x + width / 2,
        y + height / 2,
        label,
        ha="center",
        va="center",
        fontsize=8.5,
        fontweight="bold",
        color="#1f2933",
    )
    return height


def draw_decision(axes, x: float, y: float, label: str, *, size: float = 6.0) -> float:
    """Draw a decision diamond centred on ``(x, y)``."""
    diamond = Polygon(
        [
            (x, y + size / 2),
            (x + size / 2, y),
            (x, y - size / 2),
            (x - size / 2, y),
        ],
        closed=True,
        linewidth=LINE_WIDTH,
        edgecolor=DECISION_EDGE,
        facecolor=DECISION_FILL,
    )
    axes.add_patch(diamond)
    axes.text(
        x,
        y,
        label,
        ha="center",
        va="center",
        fontsize=8,
        fontweight="bold",
        color="#1f2933",
    )
    return size


def draw_arrow(
    axes,
    start: tuple[float, float],
    end: tuple[float, float],
    label: str | None = None,
    *,
    style: str = "-|>",
    color: str = "#40505f",
    rad: float = 0.0,
    dashed: bool = False,
    label_offset: tuple[float, float] = (0.0, 1.0),
) -> None:
    """Draw an arrow between two points, optionally with a label."""
    axes.annotate(
        "",
        xy=end,
        xytext=start,
        arrowprops={
            "arrowstyle": style,
            "color": color,
            "linewidth": 1.2,
            "linestyle": "--" if dashed else "-",
            "connectionstyle": f"arc3,rad={rad}",
            "shrinkA": 1,
            "shrinkB": 1,
        },
    )
    if label:
        middle_x = (start[0] + end[0]) / 2 + label_offset[0]
        middle_y = (start[1] + end[1]) / 2 + label_offset[1]
        axes.text(
            middle_x,
            middle_y,
            label,
            ha="center",
            va="bottom",
            fontsize=7.5,
            color="#33414e",
            bbox={
                "boxstyle": "round,pad=0.2",
                "facecolor": "white",
                "edgecolor": "none",
                "alpha": 0.85,
            },
        )


def draw_self_message(
    axes,
    x: float,
    y: float,
    label: str,
    *,
    reach: float = 8.0,
    drop: float = 4.0,
) -> None:
    """Draw a self message (a message an object sends to itself)."""
    axes.plot(
        [x, x + reach, x + reach, x],
        [y, y, y - drop, y - drop],
        color="#40505f",
        linewidth=1.2,
    )
    axes.annotate(
        "",
        xy=(x, y - drop),
        xytext=(x + reach, y - drop),
        arrowprops={
            "arrowstyle": "-|>",
            "color": "#40505f",
            "linewidth": 1.2,
        },
    )
    axes.text(
        x + reach / 2,
        y + 0.8,
        label,
        ha="center",
        va="bottom",
        fontsize=7.5,
        color="#33414e",
        bbox={
            "boxstyle": "round,pad=0.2",
            "facecolor": "white",
            "edgecolor": "none",
            "alpha": 0.85,
        },
    )


def draw_group(
    axes,
    x: float,
    y: float,
    width: float,
    height: float,
    label: str,
    *,
    fill: str = GROUP_FILL,
    edge: str = GROUP_EDGE,
) -> None:
    """Draw a dashed grouping frame with a caption."""
    patch = FancyBboxPatch(
        (x, y),
        width,
        height,
        boxstyle="round,pad=0.6,rounding_size=1.2",
        linewidth=1.2,
        edgecolor=edge,
        facecolor=fill,
        linestyle="--",
        alpha=0.55,
        zorder=-1,
    )
    axes.add_patch(patch)
    axes.text(
        x + 1.4,
        y + height - 0.8,
        label,
        ha="left",
        va="top",
        fontsize=8.5,
        fontweight="bold",
        color="#5b6874",
    )


def save(figure, path: Path) -> Path:
    """Write the figure to ``path`` and close it."""
    figure.savefig(path, dpi=200, bbox_inches="tight", facecolor="white")
    plt.close(figure)
    return path


# ---------------------------------------------------------------------------
# Diagrams
# ---------------------------------------------------------------------------
def activity_diagram(path: Path) -> Path:
    """Activity diagram of the capture processing pipeline."""
    figure, axes = new_canvas(11, 13)
    centre = 42.0

    draw_actor(axes, centre - 6.0, 93.0, "Capture arrives", width=12.0)

    steps: tuple[tuple[str, tuple[str, ...]], ...] = (
        (
            "Run YOLO26 detection",
            ("model/.../rain_ripple_yolo26_cpu", "CPU inference, conf = 0.25"),
        ),
        (
            "Extract ripple attributes",
            ("count, centre (x, y), radius", "bounding box, confidence"),
        ),
    )

    y = 82.0
    heights: list[float] = []
    for title, lines in steps:
        heights.append(draw_centered_box(axes, centre, y, title, lines, width=21.0))
        y -= heights[-1] + 4.0

    decision_y = y - 4.0
    draw_decision(axes, centre, decision_y, "Ripples found?")

    y = decision_y - 8.0
    heights.append(
        draw_centered_box(
            axes,
            centre,
            y,
            "Store features",
            ("SQLite: images + ripples", "JSON: storage/features/<name>.json"),
            width=22.0,
        )
    )
    y -= heights[-1] + 4.0
    heights.append(
        draw_centered_box(
            axes,
            centre,
            y,
            "Move capture",
            ("raw/ -> processed/",),
            width=21.0,
        )
    )
    y -= heights[-1] + 4.0
    heights.append(
        draw_centered_box(
            axes,
            centre,
            y,
            "Derive master key",
            ("averaged attributes -> RK4", "SHA-256 digest, 32/64 bit"),
            width=22.0,
        )
    )
    y -= heights[-1] + 4.0
    heights.append(
        draw_centered_box(
            axes,
            centre,
            y,
            "Derive slave keys",
            ("one per ripple -> RK4", "bound to master key (HMAC-SHA256)"),
            width=22.0,
        )
    )
    y -= heights[-1] + 4.0
    heights.append(
        draw_centered_box(
            axes,
            centre,
            y,
            "Store keys",
            ("SQLite: master_keys", "SQLite: slave_keys"),
            width=21.0,
        )
    )

    # arrows of the happy path
    top = 93.0
    bottom = 82.0
    draw_arrow(axes, (centre, top), (centre, bottom))
    previous_bottom = 82.0
    previous_height = heights[0]
    for index in range(1, len(heights)):
        current_top = previous_bottom - previous_height - 4.0
        if index == 2:
            # the decision sits between the extraction and the storage
            draw_arrow(
                axes,
                (centre, current_top),
                (centre, decision_y + 3.0),
            )
            draw_arrow(
                axes,
                (centre, decision_y - 3.0),
                (centre, current_top - 4.0 - 4.0),
            )
            previous_bottom = current_top - 8.0
            previous_height = heights[2]
            continue
        draw_arrow(axes, (centre, previous_bottom), (centre, current_top))
        previous_bottom = current_top
        previous_height = heights[index]

    # "no ripple" branch
    branch_x = 80.0
    draw_box(
        axes,
        branch_x,
        decision_y - 1.0,
        "No ripple detected",
        ("record status no_features",),
        width=18.0,
        fill=DECISION_FILL,
        edge=DECISION_EDGE,
    )
    delete_y = decision_y - 12.0
    draw_box(
        axes,
        branch_x,
        delete_y,
        "Delete capture",
        ("the capture is discarded",),
        width=18.0,
        fill=DECISION_FILL,
        edge=DECISION_EDGE,
    )
    draw_arrow(
        axes,
        (centre + 3.0, decision_y),
        (branch_x, decision_y + 1.0),
        "no",
        label_offset=(-1.0, 0.8),
    )
    draw_arrow(
        axes,
        (branch_x + 9.0, decision_y - 1.0),
        (branch_x + 9.0, delete_y + 5.0),
    )

    axes.text(
        50,
        99.5,
        "rainrs-server processing pipeline",
        ha="center",
        va="top",
        fontsize=12,
        fontweight="bold",
    )
    return save(figure, path)


def sequence_diagram(path: Path) -> Path:
    """Sequence diagram of the processing and derivation stages."""
    figure, axes = new_canvas(13, 10)

    actors = [
        (7.0, "CLI / Operator"),
        (28.0, "RippleFeature\nExtractor"),
        (50.0, "RainRippleDetector\n(YOLO26, torch CPU)"),
        (72.0, "SQLite\n(storage/rainrs.db)"),
        (93.0, "KeyDeriver\n(Runge-Kutta)"),
    ]
    top = 84.0
    bottom = 6.0
    for x, label in actors:
        draw_actor(axes, x - 6.75, top, label, width=13.5)
        axes.plot([x, x], [bottom, top], color="#9aa5b1", linewidth=1.0, linestyle="--")

    def message(
        x1: float,
        y: float,
        x2: float,
        label: str,
        *,
        dashed: bool = False,
    ) -> None:
        draw_arrow(
            axes,
            (x1, y),
            (x2, y),
            label,
            dashed=dashed,
            style="<|-" if dashed else "-|>",
            label_offset=(0.0, 1.0),
        )

    message(7.0, 76.0, 28.0, "process_directory(raw/)")
    message(28.0, 68.0, 50.0, "detect(image)")
    message(50.0, 60.0, 28.0, "DetectionResult(ripples)", dashed=True)
    message(28.0, 52.0, 72.0, "store_image(attributes)")
    draw_self_message(axes, 28.0, 44.0, "write attributes JSON", reach=12.0)
    draw_self_message(axes, 28.0, 34.0, "archive capture file", reach=12.0)

    message(7.0, 24.0, 93.0, "derive_for_all()")
    message(93.0, 16.0, 72.0, "load_image(id)")
    message(72.0, 10.0, 93.0, "RippleAttributes", dashed=True)

    axes.text(
        50,
        98.0,
        "Runtime sequence: capture -> features -> keys",
        ha="center",
        va="top",
        fontsize=12,
        fontweight="bold",
    )
    return save(figure, path)


def class_diagram(path: Path) -> Path:
    """Class diagram of the pipeline modules, grouped by layer."""
    figure, axes = new_canvas(14, 11)

    # -- the three layers ------------------------------------------------
    draw_group(axes, 1.0, 58.0, 27.0, 38.0, "Input layer")
    draw_group(axes, 30.0, 58.0, 39.0, 38.0, "Process layer")
    draw_group(axes, 71.0, 58.0, 28.0, 38.0, "Output layer")
    draw_group(axes, 1.0, 4.0, 98.0, 50.0, "Data model (rainrs_server.model)")

    draw_box(
        axes,
        2.0,
        70.0,
        "CaptureLibrary",
        (
            "- raw_dir: Path",
            "- processed_dir: Path",
            "+ list_captures() -> list[Path]",
            "+ archive(path) -> Path",
            "+ discard(path)",
        ),
        width=25.0,
    )
    draw_box(
        axes,
        2.0,
        61.0,
        "Image suffix rules",
        ("+ is_supported_image(path) -> bool",),
        width=25.0,
    )

    draw_box(
        axes,
        31.0,
        74.0,
        "RippleFeatureExtractor",
        (
            "- detector: RippleDetector",
            "- captures: CaptureLibrary",
            "- writer: AttributeWriter",
            "+ extract(path) -> RippleAttributes",
            "+ process_image(path) -> outcome",
        ),
        width=19.0,
        body_size=7,
    )
    draw_box(
        axes,
        31.0,
        62.0,
        "RippleWaveSystem",
        (
            "- parameters: RippleWaveParameters",
            "+ __call__(t, state) -> Vector",
        ),
        width=19.0,
        body_size=7,
    )
    draw_box(
        axes,
        54.0,
        74.0,
        "RainRippleDetector",
        (
            "- weights_path: Path",
            "- settings: Detection",
            "+ detect(path) -> result",
        ),
        width=14.0,
        body_size=7,
    )
    draw_box(
        axes,
        54.0,
        62.0,
        "KeyDeriver",
        (
            "- database_path: Path",
            "- bits: int",
            "+ derive_for_image(id)",
            "+ derive_for_all()",
        ),
        width=14.0,
        body_size=7,
    )

    draw_box(
        axes,
        72.0,
        72.0,
        "AttributeWriter",
        (
            "- features_dir: Path",
            "+ write(attributes) -> Path",
        ),
        width=26.0,
    )
    draw_box(
        axes,
        72.0,
        60.0,
        "SQLite persistence",
        (
            "open_database(path)",
            "store_image, load_image",
            "store_master_key, store_slave_key",
        ),
        width=26.0,
    )

    # -- the data model --------------------------------------------------
    model_boxes = {
        "RippleDetection": (
            3.0,
            40.0,
            ("index, center_x, center_y, radius", "width, height, confidence"),
        ),
        "RippleAttributes": (
            27.0,
            40.0,
            (
                "source_path, image_width/height",
                "ripples: tuple[RippleDetection]",
                "+ count, mean_radius, with_storage",
            ),
        ),
        "MasterKey": (
            51.0,
            40.0,
            ("image_id, key, bits, algorithm", "parameters, derived_at"),
        ),
        "Schema catalogue": (
            75.0,
            40.0,
            ("model/schema.py", "field names, types, meaning"),
        ),
        "DetectionResult": (
            3.0,
            26.0,
            ("image_width, image_height", "ripples: tuple[RippleDetection]"),
        ),
        "RippleWaveParameters": (
            27.0,
            26.0,
            ("mu, coupling, forcing, phase",),
        ),
        "SlaveKey": (
            51.0,
            26.0,
            ("master_key_id, image_id", "ripple_index, key, bits"),
        ),
        "ImageOutcome": (
            75.0,
            26.0,
            ("source, status, image_id", "ripple_count, features_path, error"),
        ),
    }
    for title, (x, y, lines) in model_boxes.items():
        draw_box(axes, x, y, title, lines, width=22.0)

    # -- relations --------------------------------------------------------
    draw_arrow(
        axes,
        (27.0, 76.0),
        (31.0, 76.0),
        "captures",
        label_offset=(0.0, 0.8),
    )
    draw_arrow(
        axes,
        (50.0, 78.0),
        (54.0, 78.0),
        "detect",
        label_offset=(0.0, 0.8),
    )
    draw_arrow(
        axes,
        (54.0, 66.0),
        (50.0, 66.0),
        "RK4",
        label_offset=(0.0, 0.8),
    )
    draw_arrow(
        axes,
        (14.0, 31.7),
        (14.0, 40.0),
        "carries",
        label_offset=(-7.0, 0.0),
    )
    draw_arrow(
        axes,
        (38.0, 31.7),
        (38.0, 40.0),
        "parameters of",
        label_offset=(8.0, 0.0),
    )
    draw_arrow(
        axes,
        (62.0, 31.7),
        (62.0, 40.0),
        "bound to",
        label_offset=(7.0, 0.0),
    )

    axes.text(
        50,
        99.0,
        "Class diagram of the rainrs-server pipeline",
        ha="center",
        va="top",
        fontsize=12,
        fontweight="bold",
    )
    return save(figure, path)


def _entity_lines(name: str, max_chars: int = 54) -> tuple[str, ...]:
    """Return the box lines of the catalogue entity called ``name``."""
    entry = data_model.entity(name)
    parts = [f"{item.name}: {item.type}" for item in entry.fields]
    lines: list[str] = []
    current = ""
    for part in parts:
        candidate = f"{current}, {part}" if current else part
        if current and len(candidate) > max_chars:
            lines.append(current)
            current = part
        else:
            current = candidate
    if current:
        lines.append(current)
    return tuple(lines)


def data_model_diagram(path: Path) -> Path:
    """Diagram of the data model, generated from the schema catalogue."""
    figure, axes = new_canvas(14, 10)

    columns = {1: 3.0, 2: 37.0, 3: 71.0}
    width = 26.0
    row_one = 66.0
    row_two = 36.0
    row_three = 6.0

    draw_group(axes, 1.0, 2.0, 98.0, 94.0, "Data model")

    draw_box(
        axes,
        columns[1],
        row_one,
        "RippleDetection",
        _entity_lines("RippleDetection"),
        width=width,
    )
    draw_box(
        axes,
        columns[2],
        row_one,
        "RippleAttributes",
        _entity_lines("RippleAttributes"),
        width=width,
    )
    draw_box(
        axes,
        columns[3],
        row_one,
        "MasterKey",
        _entity_lines("MasterKey"),
        width=width,
    )
    draw_box(
        axes,
        columns[1],
        row_two,
        "DetectionResult",
        _entity_lines("DetectionResult"),
        width=width,
    )
    draw_box(
        axes,
        columns[2],
        row_two,
        "ImageOutcome",
        _entity_lines("ImageOutcome"),
        width=width,
    )
    draw_box(
        axes,
        columns[3],
        row_two,
        "SlaveKey",
        _entity_lines("SlaveKey"),
        width=width,
    )
    draw_box(
        axes,
        columns[1],
        row_three,
        "RippleWaveParameters",
        _entity_lines("RippleWaveParameters"),
        width=width,
    )
    draw_box(
        axes,
        columns[2],
        row_three,
        "Schema catalogue",
        (
            "model/schema.py",
            "name, type and meaning of every field",
            "verified against the dataclasses",
        ),
        width=width,
    )

    # relations
    draw_arrow(
        axes,
        (columns[2], row_one + 8.0),
        (columns[1] + width, row_one + 8.0),
        "1..n ripples",
        label_offset=(0.0, 0.8),
    )
    draw_arrow(
        axes,
        (columns[2] + width, row_one + 4.0),
        (columns[3], row_one + 4.0),
        "derives",
        label_offset=(0.0, 0.8),
    )
    draw_arrow(
        axes,
        (columns[1] + 13.0, row_two + 8.0),
        (columns[1] + 13.0, row_one),
        "carries",
        label_offset=(-8.0, 0.0),
    )
    draw_arrow(
        axes,
        (columns[2] + 13.0, row_two + 8.0),
        (columns[2] + 13.0, row_one),
        "reports",
        label_offset=(9.0, 0.0),
    )
    draw_arrow(
        axes,
        (columns[3] + 13.0, row_two + 8.0),
        (columns[3] + 13.0, row_one),
        "bound to",
        label_offset=(8.0, 0.0),
    )
    draw_arrow(
        axes,
        (columns[1] + 13.0, row_three + 8.0),
        (columns[1] + 13.0, row_two),
        "parameters",
        label_offset=(-9.0, 0.0),
    )

    axes.text(
        50,
        99.5,
        "Data model of the rainrs-server pipeline",
        ha="center",
        va="top",
        fontsize=12,
        fontweight="bold",
    )
    return save(figure, path)


def erd_diagram(path: Path) -> Path:
    """Entity relationship diagram of the SQLite database."""
    figure, axes = new_canvas(13, 9)

    images_lines = (
        "id            INTEGER  PK",
        "source_path   TEXT",
        "stored_path   TEXT",
        "image_width   INTEGER",
        "image_height  INTEGER",
        "ripple_count  INTEGER",
        "status        TEXT",
        "error         TEXT",
        "features_json TEXT",
        "processed_at  TEXT",
    )
    ripples_lines = (
        "id           INTEGER  PK",
        "image_id     INTEGER  FK",
        "ripple_index INTEGER",
        "center_x     REAL",
        "center_y     REAL",
        "radius       REAL",
        "width        REAL",
        "height       REAL",
        "confidence   REAL",
    )
    master_lines = (
        "id              INTEGER  PK",
        "image_id        INTEGER  FK (unique)",
        "master_key      TEXT",
        "bits            INTEGER",
        "algorithm       TEXT",
        "parameters_json TEXT",
        "derived_at      TEXT",
    )
    slave_lines = (
        "id            INTEGER  PK",
        "master_key_id INTEGER  FK",
        "image_id      INTEGER  FK",
        "ripple_index  INTEGER",
        "slave_key     TEXT",
        "bits          INTEGER",
        "algorithm     TEXT",
        "derived_at    TEXT",
    )

    images_top = 62.0
    ripples_top = 26.0
    images_height = draw_box(axes, 3.0, images_top, "images", images_lines, width=25.0)
    ripples_height = draw_box(
        axes, 3.0, ripples_top, "ripples", ripples_lines, width=25.0
    )
    draw_box(axes, 62.0, images_top, "master_keys", master_lines, width=25.0)
    slave_height = draw_box(
        axes, 62.0, ripples_top, "slave_keys", slave_lines, width=25.0
    )

    draw_arrow(
        axes,
        (15.5, images_top),
        (15.5, ripples_top + ripples_height),
        "1 : N",
        label_offset=(2.6, 0.0),
    )
    draw_arrow(
        axes,
        (28.0, images_top + images_height - 3.0),
        (62.0, images_top + images_height - 3.0),
        "1 : 1",
        label_offset=(0.0, 1.0),
    )
    draw_arrow(
        axes,
        (74.5, images_top),
        (74.5, ripples_top + slave_height),
        "1 : N",
        label_offset=(2.6, 0.0),
    )
    draw_arrow(
        axes,
        (28.0, ripples_top + ripples_height - 3.0),
        (62.0, ripples_top + slave_height - 3.0),
        "1 : N",
        label_offset=(0.0, 1.0),
    )

    axes.text(
        50,
        97.0,
        "Entity relationship diagram - storage/rainrs.db",
        ha="center",
        va="top",
        fontsize=12,
        fontweight="bold",
    )
    axes.text(
        50,
        6.0,
        "images -> ripples (cascade delete)    |    "
        "images -> master_keys (1:1)    |    "
        "master_keys -> slave_keys (cascade delete)",
        ha="center",
        va="bottom",
        fontsize=7.5,
        color="#33414e",
    )
    return save(figure, path)


def component_diagram(path: Path) -> Path:
    """Component diagram of the server and its environment."""
    figure, axes = new_canvas(13, 9)

    inner_top = 52.0
    inner_bottom = 4.0
    draw_group(axes, 2.0, inner_top, 96.0, 44.0, "rainrs-server")
    draw_group(axes, 2.0, inner_bottom, 96.0, 44.0, "Environment")

    draw_box(
        axes,
        4.0,
        64.0,
        "CLI (cli.py)",
        ("python -m rainrs_server", "process | derive | run"),
        width=20.0,
    )
    draw_box(
        axes,
        28.0,
        64.0,
        "Input layer",
        ("input/captures.py", "CaptureLibrary"),
        width=20.0,
    )
    draw_box(
        axes,
        52.0,
        64.0,
        "Process layer",
        ("process/image_read.py", "process/rungge.py"),
        width=20.0,
    )
    draw_box(
        axes,
        76.0,
        64.0,
        "Output layer",
        ("output/attributes.py", "output/database.py"),
        width=20.0,
    )

    draw_box(
        axes,
        4.0,
        10.0,
        "Capture sources",
        ("desktop / mobile app", "embedded system (ESP32)"),
        width=20.0,
        fill=ACTOR_FILL,
        edge=ACTOR_EDGE,
    )
    draw_box(
        axes,
        28.0,
        10.0,
        "File system",
        ("storage/images/raw", "storage/images/processed", "storage/features"),
        width=20.0,
        fill=STORE_FILL,
        edge=STORE_EDGE,
    )
    draw_box(
        axes,
        52.0,
        10.0,
        "YOLO26 model",
        ("model/.../best.pt", "ultralytics + torch (CPU)"),
        width=20.0,
        fill=STORE_FILL,
        edge=STORE_EDGE,
    )
    draw_box(
        axes,
        76.0,
        10.0,
        "Logging",
        ("logs/rainrs.log", "no features or keys"),
        width=20.0,
        fill=STORE_FILL,
        edge=STORE_EDGE,
    )

    draw_arrow(axes, (14.0, 20.0), (14.0, 64.0), "captures")
    draw_arrow(axes, (24.0, 70.0), (28.0, 70.0))
    draw_arrow(axes, (48.0, 70.0), (52.0, 70.0))
    draw_arrow(axes, (72.0, 70.0), (76.0, 70.0))
    draw_arrow(axes, (38.0, 64.0), (38.0, 20.0), "reads / archives")
    draw_arrow(axes, (62.0, 64.0), (62.0, 20.0), "loads weights", dashed=True)
    draw_arrow(axes, (86.0, 64.0), (86.0, 20.0), "writes", dashed=True)

    axes.text(
        50,
        98.0,
        "Component diagram of rainrs-server",
        ha="center",
        va="top",
        fontsize=12,
        fontweight="bold",
    )
    return save(figure, path)


# ---------------------------------------------------------------------------
# Document helpers
# ---------------------------------------------------------------------------
def add_code(document: Document, code: str) -> None:
    """Append a monospaced code block."""
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    run = paragraph.add_run(code)
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def add_table(
    document: Document, header: tuple[str, ...], rows: tuple[tuple[str, ...], ...]
) -> None:
    """Append a simple table with a header row."""
    table = document.add_table(rows=1, cols=len(header))
    table.style = "Light Grid Accent 1"
    for cell, title in zip(table.rows[0].cells, header, strict=True):
        cell.text = title
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            cell.text = value


def build_document(diagrams: dict[str, Path]) -> Document:
    """Assemble the documentation document."""
    document = Document()
    document.core_properties.title = "rainrs-server documentation"
    document.core_properties.author = "Joshua Elijah L. Pura"
    document.core_properties.subject = "Rain Ripple Key Generator - server component"

    document.add_heading("rainrs-server documentation", level=0)
    subtitle = document.add_paragraph(
        "Rain Ripple Key Generator - server component\n"
        "Capture processing, ripple feature extraction and Runge-Kutta key derivation"
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(
        "Version 0.1.0 | 22 October 2026 | "
        "Repository: Fauzst/RainRippleV3 | Branch: feat/image_process"
    )

    document.add_heading("Table of contents", level=1)
    for entry in (
        "1. Overview",
        "2. Component diagram",
        "3. Activity diagram (processing pipeline)",
        "4. Class diagram (module design)",
        "5. Sequence diagram (runtime behaviour)",
        "6. Entity relationship diagram (database)",
        "7. Data model",
        "8. Data flow and storage layout",
        "9. Command line usage",
        "10. Security considerations",
        "11. Limitations and known issues",
        "12. Verification results",
    ):
        document.add_paragraph(entry, style="List Bullet")

    document.add_heading("1. Overview", level=1)
    document.add_paragraph(
        "The server component of the Rain Ripple Key Generator turns captured rain "
        "ripple images into experimental key material. The pipeline has two stages."
    )
    document.add_paragraph(
        "Stage 1 - feature extraction: every capture in storage/images/raw is passed "
        "through the trained YOLO26 detector "
        "(model/content/runs/detect/rain_ripple_yolo26_cpu/weights/best.pt). For every "
        "detected ripple the centre position, the bounding box and the equivalent "
        "radius are measured. The measurements are stored in the SQLite database "
        "and in a JSON document (storage/features/<name>.json). Captures with "
        "measurable ripples are moved to storage/images/processed; captures without "
        "any ripple are deleted."
    )
    document.add_paragraph(
        "Stage 2 - key derivation: the stored attributes are fed into a Runge-Kutta "
        "integration of a coupled oscillator system. The averaged attributes of an "
        "image produce one master key; every individual ripple produces one slave "
        "key that is bound to its master key. All keys are written to the SQLite "
        "database."
    )
    document.add_paragraph(
        "The application is organised as a simple MVC split. The model "
        "(rainrs_server.model) describes every value that flows through the "
        "pipeline; the input layer (rainrs_server.input) owns the incoming captures; "
        "the process layer (rainrs_server.process) computes with the model; and the "
        "output layer (rainrs_server.output) stores the results. Section 7 shows the "
        "shape of the data itself."
    )
    document.add_paragraph(
        "The generated material is an experimental research artifact. It is not a "
        "production cryptographic key - see section 10."
    )

    document.add_heading("2. Component diagram", level=1)
    document.add_picture(str(diagrams["component"]), width=Inches(6.5))
    caption = document.add_paragraph(
        "Figure 1 - Components of rainrs-server and their environment."
    )
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading("3. Activity diagram (processing pipeline)", level=1)
    document.add_picture(str(diagrams["activity"]), width=Inches(5.6))
    caption = document.add_paragraph(
        "Figure 2 - Activity flow of the capture processing pipeline."
    )
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading("4. Class diagram (module design)", level=1)
    document.add_picture(str(diagrams["class"]), width=Inches(6.5))
    caption = document.add_paragraph(
        "Figure 3 - Classes of the input layer, the process layer, the output layer "
        "and the data model."
    )
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading("5. Sequence diagram (runtime behaviour)", level=1)
    document.add_picture(str(diagrams["sequence"]), width=Inches(6.5))
    caption = document.add_paragraph(
        "Figure 4 - Runtime sequence between the CLI, the extraction stage, the "
        "detector, the database and the derivation stage. The extractor reads the "
        "captures through the input layer (CaptureLibrary) and writes through the "
        "output layer (AttributeWriter and the SQLite database)."
    )
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading("6. Entity relationship diagram (database)", level=1)
    document.add_picture(str(diagrams["erd"]), width=Inches(6.5))
    caption = document.add_paragraph(
        "Figure 5 - Tables of storage/rainrs.db and their relationships."
    )
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_table(
        document,
        ("Table", "Purpose", "Key columns"),
        (
            (
                "images",
                "One row per processed capture",
                "id, source_path, stored_path, status, features_json",
            ),
            (
                "ripples",
                "One row per detected ripple",
                "id, image_id (FK), center_x/y, radius, confidence",
            ),
            (
                "master_keys",
                "Exactly one master key per capture",
                "id, image_id (FK, unique), master_key, bits, algorithm",
            ),
            (
                "slave_keys",
                "One row per ripple, bound to a master key",
                "id, master_key_id (FK), image_id (FK), ripple_index, slave_key",
            ),
        ),
    )

    document.add_heading("7. Data model", level=1)
    document.add_paragraph(
        "Every value that flows through the pipeline is described in the model "
        "package (rainrs_server/model). The package holds the dataclasses of the "
        "data and a declarative catalogue (rainrs_server/model/schema.py) that names "
        "the field, the type and the meaning of every entity. The catalogue is not a "
        "copy that can drift away from the code: the test suite compares it with the "
        "dataclasses and fails when a field was added, renamed or retyped without "
        "updating the catalogue."
    )
    document.add_picture(str(diagrams["data_model"]), width=Inches(6.5))
    caption = document.add_paragraph(
        "Figure 6 - The data model of the pipeline, generated from the catalogue in "
        "rainrs_server/model/schema.py."
    )
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(
        "The catalogue covers the following entities. The 'Stored in' column names "
        "the SQLite table the entity is persisted in; entities without a table are "
        "transient values that only live inside the pipeline."
    )
    add_table(
        document,
        ("Entity", "Stored in", "Purpose"),
        tuple(
            (entry.name, entry.table or "-", entry.description)
            for entry in data_model.entities()
        ),
    )
    document.add_paragraph("The fields of every entity:")
    for entry in data_model.entities():
        document.add_heading(entry.name, level=2)
        document.add_paragraph(entry.description)
        add_table(
            document,
            ("Field", "Type", "Nullable", "Description"),
            tuple(
                (
                    item.name,
                    item.type,
                    "yes" if item.nullable else "no",
                    item.description,
                )
                for item in entry.fields
            ),
        )

    document.add_heading("8. Data flow and storage layout", level=1)
    document.add_paragraph(
        "The application package follows the MVC split; the runtime artefacts live "
        "next to it:"
    )
    add_code(
        document,
        "rainrs-server/\n"
        "  src/rainrs_server/\n"
        "    model/      the data model and its catalogue (schema.py)\n"
        "    input/      the capture library (captures.py)\n"
        "    process/    detection (image_read.py) and derivation (rungge.py)\n"
        "    output/     attribute documents (attributes.py),\n"
        "                SQLite persistence (database.py)\n"
        "  model/content/runs/detect/rain_ripple_yolo26_cpu/weights/best.pt\n"
        "  storage/\n"
        "    images/raw/          incoming captures (deleted when no ripple)\n"
        "    images/processed/    captures whose features were extracted\n"
        "    features/            JSON attribute documents\n"
        "    rainrs.db            SQLite database (features + keys)\n"
        "  logs/rainrs.log       pipeline log (no features, no key material)\n"
        "  docs/                 this document",
    )
    document.add_paragraph("The attribute document of one capture looks like this:")
    add_code(
        document,
        "{\n"
        '  "image_id": 3,\n'
        '  "image_width": 200,\n'
        '  "image_height": 113,\n'
        '  "ripple_count": 2,\n'
        '  "mean_center_x": 95.01714324951172,\n'
        '  "mean_center_y": 55.14802551269531,\n'
        '  "mean_radius": 10.720739364624023,\n'
        '  "ripples": [\n'
        '    {"index": 0, "center_x": 118.47, "center_y": 37.17,\n'
        '     "radius": 15.68, "width": 29.61, "height": 33.13,\n'
        '     "confidence": 0.5669}\n'
        "  ]\n"
        "}",
    )

    document.add_heading("9. Command line usage", level=1)
    document.add_paragraph(
        "All commands are executed from the rainrs-server directory with uv:"
    )
    add_code(
        document,
        "uv sync                       # install the dependencies\n"
        "uv run python -m rainrs_server process --print\n"
        "uv run python -m rainrs_server derive --bits 64\n"
        "uv run python -m rainrs_server run --print",
    )
    add_table(
        document,
        ("Command", "Effect"),
        (
            (
                "process",
                "Detects ripples in storage/images/raw and stores the features.",
            ),
            (
                "derive",
                "Derives the master and slave keys from the stored attributes.",
            ),
            (
                "run",
                "Runs the detection stage and the derivation stage in sequence.",
            ),
            ("--print", "Enables the development console printing (stdout only)."),
            ("--no-move", "Keeps the captures in the raw directory (dry run)."),
            ("--keep-no-features", "Keeps captures without a detected ripple."),
            ("--bits 32|64", "Size of the generated research artifacts."),
            ("--binding hmac|xor", "Binding of the slave keys to their master key."),
        ),
    )
    document.add_paragraph(
        "The development printing is switched on with the environment variable "
        "RAINRS_DEV_PRINT=1 and is implemented in a single place "
        "(rainrs_server/process/dev.py), so it can be removed with one edit. It "
        "writes to stdout only and never to the log file."
    )

    document.add_heading("10. Security considerations", level=1)
    for note in (
        "The generated 32/64 bit values are experimental research artifacts. They "
        "are far too short for cryptographic use and have not been evaluated with "
        "a statistical test suite (for example NIST SP 800-22).",
        "The derivation is deterministic: the same capture always produces the same "
        "key. The unpredictability of the artifact comes from the physical capture, "
        "not from the algorithm. An attacker who knows the capture can recompute "
        "the key.",
        "Production encryption must use established cryptographic libraries, "
        "approved algorithms and keys generated by a CSPRNG.",
        "The slave keys are bound to their master key with HMAC-SHA256 instead of a "
        "plain XOR: XOR lets two equal inputs cancel out and does not separate the "
        "child key from the parent. The XOR variant remains available for "
        "comparison.",
        "The log file never contains captures, derived features or key material. "
        "Only file names, counters, statuses and timings are recorded.",
        "The SQLite database holds derived features and key material; it is "
        "excluded from version control.",
    ):
        document.add_paragraph(note, style="List Bullet")

    document.add_heading("11. Limitations and known issues", level=1)
    for note in (
        "The detector runs on the CPU; a single 200x113 pixel capture needs a few "
        "seconds, so the pipeline is not suited for a high capture rate.",
        "The trained model was trained at 416 pixel inputs; the captures of the "
        "sample set are much smaller and are upscaled during inference.",
        "The key derivation has not been evaluated statistically. No entropy "
        "estimate, no min-entropy analysis and no statistical test results exist "
        "yet.",
        "The extraction stage stores the ripple geometry in pixels of the source "
        "image. A capture with a different resolution therefore produces different "
        "attributes, which is intended but has to be kept in mind when comparing "
        "captures.",
        "The CLI writes the derived key material only to the database; there is no "
        "API surface yet. The API design is still undecided in agent/PLAN.md.",
    ):
        document.add_paragraph(note, style="List Bullet")

    document.add_heading("12. Verification results", level=1)
    document.add_paragraph(
        "The pipeline was executed on the ten sample captures of "
        "storage/images/raw on the reference host (Debian 13, Intel Core i5-6500, "
        "4 cores, no GPU). Every capture was processed successfully."
    )
    add_table(
        document,
        ("Stage", "Result"),
        (
            ("Detection", "10 of 10 captures processed, 0 errors, 37 ripples detected"),
            (
                "No-feature branch",
                "verified with conf = 0.95: capture deleted, status no_features",
            ),
            (
                "Error branch",
                "verified with a corrupt capture: error logged, capture kept",
            ),
            ("Key derivation", "10 master keys and 37 slave keys stored"),
            (
                "Data model",
                "the catalogue of rainrs_server/model/schema.py matches the "
                "dataclasses (verified by the test suite)",
            ),
            ("Unit tests", "79 tests pass (pytest)"),
            ("Lint / format", "ruff check and ruff format --check pass"),
            ("Type check", "mypy reports no issues"),
        ),
    )
    document.add_paragraph(
        "The measured numbers are documented in agent/PROGRESS.md of the "
        "rainrs-server application."
    )

    return document


def main() -> int:
    """Render the diagrams and write the documentation."""
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="rainrs-docs-") as workdir:
        target = Path(workdir)
        diagrams = {
            "activity": activity_diagram(target / "activity.png"),
            "sequence": sequence_diagram(target / "sequence.png"),
            "class": class_diagram(target / "class.png"),
            "erd": erd_diagram(target / "erd.png"),
            "component": component_diagram(target / "component.png"),
            "data_model": data_model_diagram(target / "data_model.png"),
        }
        document = build_document(diagrams)
        document.save(OUTPUT)
    print(f"documentation written to {OUTPUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
